#!/usr/bin/env python3
"""render.py — build derived artifacts from the store.

- build_payload(store_dict, cfg): the dashboard JSON payload
- write_dashboard(...): career_data.json (local plaintext) + data.enc.json
  (AES-GCM encrypted blob for the public/phone deployment)
- render_xlsx(...): the Job Tracker.xlsx spreadsheet mirror (openpyxl)
"""
import json
import os
from datetime import datetime, timezone

from engine import analytics, classify, crypto, recommend

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
WEBDIR = os.path.join(ROOT, "web")
PLAIN = os.path.join(WEBDIR, "career_data.json")
ENC = os.path.join(WEBDIR, "data.enc.json")


def build_payload(store, cfg, sample=False):
    """Compute everything the dashboard needs from the raw store dict."""
    apps = analytics.enrich(store, cfg)
    metrics = analytics.compute_metrics(apps)
    recs = recommend.recommendations(apps, store, metrics, cfg)
    templates = [{"name": t.get("name", ""), "audience": t.get("audience", ""),
                  "when": t.get("when", t.get("when_to_use", "")),
                  "body": t.get("body", "")}
                 for t in store.get("templates", [])]
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "owner": store.get("meta", {}).get("owner_name", cfg.get("owner_name", "You")),
        "sample": sample,
        "stage_order": cfg["stages"]["order"] + cfg["stages"].get("terminal", []),
        "metrics": metrics,
        "recommendations": recs,
        "applications": sorted(apps, key=lambda a: a["id"]),
        "interviews": sorted(store.get("interviews", []),
                             key=lambda x: x.get("date", ""), reverse=True),
        "contacts": store.get("contacts", []),
        "networking_queue": analytics.networking_queue(store, cfg),
        "templates": templates,
        "leads": sorted(store.get("leads", []),
                        key=lambda l: l.get("date", ""), reverse=True)[:60],
    }


def write_dashboard(store, cfg, sample=False, webdir=None):
    """Write career_data.json + data.enc.json into webdir. Returns the payload."""
    webdir = webdir or WEBDIR
    os.makedirs(webdir, exist_ok=True)
    payload = build_payload(store, cfg, sample=sample)
    plain_path = os.path.join(webdir, "career_data.json")
    with open(plain_path, "w") as f:
        json.dump(payload, f, indent=2)
    enc_path = os.path.join(webdir, "data.enc.json")
    try:
        blob = crypto.encrypt_payload(payload, crypto.get_passphrase())
        with open(enc_path, "w") as f:
            json.dump(blob, f)
        enc_msg = f"encrypted → {enc_path}"
    except ImportError:
        enc_msg = ("encryption skipped: 'cryptography' not installed — "
                   "local dashboard still works")
    m = payload["metrics"]
    print(f"build OK — {m['total']} apps, {m['interview_rate']}% interview rate, "
          f"{len(payload['recommendations'])} recommendations, "
          f"{len(payload['networking_queue'])} contacts queued.")
    print(f"  plaintext → {plain_path}\n  {enc_msg}")
    return payload


def render_xlsx(store, cfg, out_path=None):
    """Render the Job Tracker.xlsx mirror from the store (openpyxl)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    apps = analytics.enrich(store, cfg)
    metrics = analytics.compute_metrics(apps)
    out_path = out_path or os.path.join(ROOT, "Job Tracker.xlsx")

    HFILL = PatternFill("solid", fgColor="1F4E79")
    HF = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    YEL = PatternFill("solid", fgColor="FFFF00")
    DF = Font(name="Arial", size=10)
    BF = Font(name="Arial", size=10, bold=True)
    C = Alignment("center", "center", wrap_text=True)
    L = Alignment("left", "center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    BD = Border(thin, thin, thin, thin)
    STAGE_COLOR = {"Rejected": "FF0000", "Final Round": "2E7D32",
                   "Round 2": "2E7D32", "Offer": "2E7D32",
                   "Interview": "1565C0", "Phone Screen": "E65100"}

    wb = Workbook()

    def sheet(name, headers, widths):
        ws = wb.create_sheet(name)
        for i, h in enumerate(headers, 1):
            c = ws.cell(1, i, h)
            c.fill = HFILL; c.font = HF; c.alignment = C; c.border = BD
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)

    # Application Tracker
    a = sheet("Application Tracker",
              ["App #", "Job Title", "Company", "Job Desc. (hyperlink)",
               "Application Stage", "Date Applied", "Source",
               "People to reach out to (LinkedIn)", "Outreach Message",
               "Reply Email Received", "Interview Received?"],
              [6, 40, 24, 20, 18, 13, 16, 26, 26, 18, 16])
    for r, ap in enumerate(sorted(apps, key=lambda x: x["id"]), 2):
        fill = YEL if ap.get("flag") else None
        vals = [ap["id"], ap["title"], ap["company"], ap.get("url", ""),
                ap["stage"], ap["date_applied"], ap["source"], "", "",
                "Yes" if ap.get("reply") else "",
                "Yes" if ap.get("interview") else "No"]
        for i, v in enumerate(vals, 1):
            c = a.cell(r, i, v)
            c.font = BF if i == 1 else DF
            c.border = BD
            c.alignment = C if i in (1, 5, 6, 7, 10, 11) else L
            if fill:
                c.fill = fill
            if i == 5 and ap["stage"] in STAGE_COLOR:
                c.font = Font(name="Arial", size=10, bold=True,
                              color=STAGE_COLOR[ap["stage"]])
            if i == 11 and ap.get("interview"):
                c.font = Font(name="Arial", size=10, bold=True, color="2E7D32")

    # Interview Tracker
    it = sheet("Interview Tracker",
               ["Interview #", "Company name", "Job Title", "Interview Date",
                "Interview Round", "Interviewer Details", "Interview Details"],
               [11, 24, 38, 14, 30, 52, 58])
    for r, iv in enumerate(sorted(store.get("interviews", []),
                                  key=lambda x: x.get("date", "")), 2):
        vals = [iv["id"], iv["company"], iv["title"], iv["date"],
                iv["round"], iv["interviewer"], iv["details"]]
        for i, v in enumerate(vals, 1):
            c = it.cell(r, i, v)
            c.font = BF if i == 1 else DF
            c.alignment = C if i in (1, 4) else L
            c.border = BD
        it.row_dimensions[r].height = 38

    # Contacts
    ct = sheet("Contacts",
               ["Name", "Title", "Company", "Email", "Phone", "Warmth",
                "Last Touch", "Context", "LinkedIn"],
               [20, 30, 22, 30, 15, 10, 13, 50, 28])
    for r, c0 in enumerate(store.get("contacts", []), 2):
        vals = [c0.get("name"), c0.get("title"), c0.get("company"),
                c0.get("email"), c0.get("phone", ""), c0.get("warmth"),
                c0.get("last_touch", ""), c0.get("context", ""),
                c0.get("linkedin", "")]
        for i, v in enumerate(vals, 1):
            c = ct.cell(r, i, v)
            c.font = DF; c.alignment = L; c.border = BD

    # Outreach Templates
    tp = sheet("Outreach Templates", ["Template", "Audience", "When to use", "Message"],
               [34, 16, 40, 90])
    for r, t in enumerate(store.get("templates", []), 2):
        vals = [t["name"], t["audience"], t.get("when", t.get("when_to_use", "")),
                t["body"]]
        for i, v in enumerate(vals, 1):
            c = tp.cell(r, i, v)
            c.font = DF; c.alignment = L; c.border = BD
        tp.row_dimensions[r].height = 120

    # Leads (from job alerts)
    lsheet = sheet("Leads (from alerts)",
                   ["Company", "Job Title", "Likely Family", "Source", "Date"],
                   [24, 40, 22, 18, 14])
    for r, le in enumerate(sorted(store.get("leads", []),
                                  key=lambda x: x.get("date", ""), reverse=True), 2):
        vals = [le.get("company"), le.get("title"),
                classify.job_family(le.get("title", ""), cfg),
                le.get("source"), le.get("date")]
        for i, v in enumerate(vals, 1):
            c = lsheet.cell(r, i, v)
            c.font = DF; c.alignment = L if i in (1, 2, 3) else C; c.border = BD

    # Insights
    ins = sheet("Insights", ["Metric", "Value"], [34, 16])
    rows = [("Total applications", metrics["total"]),
            ("Reached interview", metrics["interviews"]),
            ("Interview rate", f'{metrics["interview_rate"]}%'),
            ("Rejected", metrics["rejected"]),
            ("Still active", metrics["active"]),
            ("Applied last 7 days", metrics["applied_last_7"]),
            ("Applied last 30 days", metrics["applied_last_30"]),
            ("", ""), ("Interview rate by family", "")]
    for fam in metrics["by_family"]:
        rows.append((f'  {fam["segment"]}',
                     f'{fam["interview_rate"]}% ({fam["interviews"]}/{fam["count"]})'))
    rows += [("", ""), ("Interview rate by source", "")]
    for s in metrics["by_source"]:
        if s["count"] >= 2:
            rows.append((f'  {s["segment"]}',
                         f'{s["interview_rate"]}% ({s["interviews"]}/{s["count"]})'))
    for r, (k, v) in enumerate(rows, 2):
        ck = ins.cell(r, 1, k); cv = ins.cell(r, 2, v)
        ck.font = BF if (k and not k.startswith("  ")) else DF
        cv.font = DF
        ck.alignment = L; cv.alignment = L; ck.border = BD; cv.border = BD

    # Config
    cf = sheet("Config", ["Key", "Value"], [22, 28])
    cursor = store.get("cursor", {})
    for r, (k, v) in enumerate([
            ("last_scan_iso", cursor.get("last_scan_iso")),
            ("processed_threads", cursor.get("processed_count")),
            ("total_applications", metrics["total"]),
            ("total_interviews", len(store.get("interviews", [])))], 2):
        ck = cf.cell(r, 1, k); cv = cf.cell(r, 2, v)
        ck.font = DF; cv.font = DF; ck.border = BD; cv.border = BD

    wb.save(out_path)
    print(f"xlsx mirror written → {out_path}")

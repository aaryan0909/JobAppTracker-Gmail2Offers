#!/usr/bin/env python3
"""
career_lib.py — the brain of the career decision board.

Commands:
  build                 Recompute everything → web/career_data.json (+ encrypted web/data.enc.json)
  cursor                Print scan cursor (last_scan_iso + processed thread-id count) as JSON
  ingest <records.json> Merge new scanned records into store.json, update cursor, then build

The store (engine/store.json) holds RAW facts. This lib computes all derived
analytics, recommendations, and the dashboard payload, then encrypts it so it
can live on a public Vercel URL while staying private.
"""
import json, os, sys, base64, re
from datetime import datetime, timezone, date

def _norm(s):
    """Normalize company/title for de-duplication: lowercase, &→and, strip punctuation."""
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower().replace("&", " and ")).strip()

def _app_key(company, title, date=""):
    """De-dup key. Placeholder/unknown titles are disambiguated by date so that
    multiple real applications with an unknown title are NOT collapsed together."""
    t = _norm(title)
    if (not t) or ("unknown" in t) or (title or "").strip().startswith("["):
        return (_norm(company), t, (date or "")[:10])
    return (_norm(company), t)

HERE   = os.path.dirname(os.path.abspath(__file__))
ROOT   = os.path.dirname(HERE)
STORE  = os.path.join(HERE, "store.json")
WEBDIR = os.path.join(ROOT, "web")
PLAIN  = os.path.join(WEBDIR, "career_data.json")
ENC    = os.path.join(WEBDIR, "data.enc.json")
PASSF  = os.path.join(HERE, ".passphrase")
TODAY  = date.today()   # dynamic; uses the machine clock

# ────────────────────────────────────────────────────────── classification
FAMILY_RULES = [
    ("Risk & Finance", ["risk","fraud","credit","collections","fp&a","financial","lending","policy","procurement","actuar","treasury"]),
    ("Data & Analytics", ["data analyst","data scientist","analytics","data steward","insights","business intelligence","bi ","reporting","metrics","statistical","modeling"]),
    ("Product", ["product manager","product owner","pm ","observability"]),
    ("Consulting", ["consultant","consulting"]),
    ("Engineering", ["engineer","developer","full stack","software"]),
    ("Operations & Strategy", ["operations","operational","strategy","gtm","go-to-market","process","settlements","continuous improvement","business operations","biz ops","bizops","onboarding","implementation","enablement","management consultant"]),
]
def job_family(title):
    t = (title or "").lower()
    # product analyst leans analytics
    if "product analyst" in t or "product & business" in t or "business intelligence analyst" in t:
        return "Data & Analytics"
    for fam, kws in FAMILY_RULES:
        if any(k in t for k in kws):
            return fam
    if "analyst" in t:
        return "Data & Analytics"
    if "associate" in t or "specialist" in t or "coordinator" in t or "representative" in t:
        return "Operations & Strategy"
    return "Other"

def seniority(title):
    t = (title or "").lower()
    if any(k in t for k in ["director","svp","vp","head of","principal","staff"]):
        return "Director+"
    if any(k in t for k in ["senior","sr.","sr ","lead","manager"]):
        return "Senior"
    if any(k in t for k in ["associate","analyst","coordinator","specialist","representative","junior","intern","co-op","steward i"]):
        return "Entry/Associate"
    return "Mid"

SECTOR = {
    "Financeit":"Fintech","RBC":"Banking","BMO":"Banking","BMO Capital Markets":"Banking",
    "Scotiabank":"Banking","TD":"Banking","American Express":"Fintech","Wealthsimple":"Fintech",
    "Ramp":"Fintech","Kraken":"Crypto","Baselayer":"Fintech","Interac":"Fintech",
    "FanDuel":"Gaming","Sporty Group":"Gaming","White Hat Gaming / Moonspin":"Gaming","Betty Gaming Canada":"Gaming",
    "DoorDash":"Tech/Marketplace","Uber":"Tech/Marketplace","Delivery Hero":"Tech/Marketplace","HelloFresh":"Tech/Marketplace",
    "OpenAI":"AI","Anthropic":"AI","Cohere":"AI","API":"AI","Samsara":"Tech","ServiceTitan":"Tech",
    "Okta":"Tech","Spotify":"Tech","AppLovin":"Tech","StackAdapt":"Tech","Roku":"Tech","Geotab":"Tech",
    "Relay":"Fintech","Dandy":"Tech","Klick":"Marketing/Health","Omnicom Media Group CA Toronto":"Marketing",
    "Viral Nation Inc.":"Marketing","Home Depot Canada":"Retail","Konrad":"Consulting","Forma.ai":"Tech",
    "Metergy Solutions":"Energy","PointClickCare":"HealthTech","MD Analytics":"Health Research",
    "Plug and Play":"VC/Startup","Leadership Connect":"Data/Media","Passport":"Logistics",
    "City of Toronto":"Public Sector","Toronto Community Housing":"Public Sector","LinkedIn":"Tech",
    "Aylo":"Tech","Canopy Planet Society":"Nonprofit","FutureFit AI":"AI","American Iron and Metal":"Industrial",
    "via Weekday Platform":"Platform","Besty":"Tech",
}

REJECTED   = {"rejected"}
def is_rejected(stage): return (stage or "").strip().lower() in REJECTED
ACTIVE_HOT = ["offer","final round","round 2","round 3","interview","phone screen","technical","onsite"]
def is_active_advanced(stage):
    s = (stage or "").lower()
    return any(k in s for k in ACTIVE_HOT)

def days_since(d):
    try:
        return (TODAY - date.fromisoformat(d[:10])).days
    except Exception:
        return None

# ────────────────────────────────────────────────────────── analytics
def enrich(store):
    apps = []
    for a in store["applications"]:
        e = dict(a)
        e["family"]   = job_family(a["title"])
        e["seniority"]= seniority(a["title"])
        e["sector"]   = SECTOR.get(a["company"], "Other")
        e["days_since"] = days_since(a.get("date_applied",""))
        e["rejected"] = is_rejected(a["stage"])
        e["advanced"] = bool(a.get("interview")) or is_active_advanced(a["stage"])
        e["stale"]    = (a["stage"].strip().lower()=="applied" and not a.get("interview")
                         and (e["days_since"] or 0) > 21 and not e["rejected"])
        apps.append(e)
    return apps

def seg_stats(apps, key):
    out = {}
    for a in apps:
        k = a[key]
        d = out.setdefault(k, {"segment":k,"count":0,"interviews":0,"replies":0,"rejected":0})
        d["count"] += 1
        d["interviews"] += 1 if a["advanced"] else 0
        d["replies"] += 1 if a.get("reply") else 0
        d["rejected"] += 1 if a["rejected"] else 0
    for d in out.values():
        d["interview_rate"] = round(100*d["interviews"]/d["count"]) if d["count"] else 0
        d["reply_rate"]     = round(100*d["replies"]/d["count"]) if d["count"] else 0
    return sorted(out.values(), key=lambda x:(-x["interview_rate"], -x["count"]))

def compute_metrics(apps):
    n = len(apps)
    interviews = sum(1 for a in apps if a["advanced"])
    replies    = sum(1 for a in apps if a.get("reply"))
    rejected   = sum(1 for a in apps if a["rejected"])
    active     = sum(1 for a in apps if not a["rejected"])
    recent7    = sum(1 for a in apps if (a["days_since"] or 999) <= 7)
    recent30   = sum(1 for a in apps if (a["days_since"] or 999) <= 30)
    return {
        "total": n, "interviews": interviews, "replies": replies,
        "rejected": rejected, "active": active,
        "interview_rate": round(100*interviews/n) if n else 0,
        "reply_rate": round(100*replies/n) if n else 0,
        "applied_last_7": recent7, "applied_last_30": recent30,
        "by_family": seg_stats(apps, "family"),
        "by_source": seg_stats(apps, "source"),
        "by_seniority": seg_stats(apps, "seniority"),
        "by_sector": seg_stats(apps, "sector"),
        "stage_dist": _stage_dist(apps),
    }

def _stage_dist(apps):
    out = {}
    for a in apps:
        out[a["stage"]] = out.get(a["stage"],0)+1
    return out

# ────────────────────────────────────────────────────────── recommendations
def recommendations(apps, store, metrics):
    recs = []
    def add(priority, kind, title, detail, action=None, items=None):
        recs.append({"priority":priority,"kind":kind,"title":title,
                     "detail":detail,"action":action,"items":items or []})

    # 1) Active advanced opportunities — highest priority
    hot = [a for a in apps if a["advanced"] and not a["rejected"]
           and a["stage"].strip().lower() not in ("applied",)]
    hot_sorted = sorted(hot, key=lambda a:(("offer" not in a["stage"].lower()), a["company"]))
    if hot_sorted:
        items=[{"label":f'{a["company"]} — {a["title"]}',"sub":a["stage"]} for a in hot_sorted[:8]]
        add(1,"opportunity","Active opportunities need your attention",
            f'{len(hot_sorted)} application(s) are mid-process. Keep momentum: confirm timing, prep, and send thank-yous within 24h of each round.',
            "Open Pipeline → focus on these", items)

    # 2) Offer decision
    offers = [a for a in apps if "offer" in a["stage"].lower()]
    if offers:
        items=[{"label":f'{a["company"]} — {a["title"]}',"sub":a["stage"]} for a in offers]
        add(1,"decision","You have an offer / imminent offer",
            "Decide and respond promptly — even a decline should be gracious and fast to protect the relationship. A decline draft is in your Templates tab.",
            "Send your decision", items)

    # 3) Best-converting segment → double down
    fam = [s for s in metrics["by_family"] if s["count"]>=3]
    if fam:
        best = fam[0]
        if best["interview_rate"]>0:
            add(2,"strategy",f'Lean into {best["segment"]} roles',
                f'Your strongest segment: {best["interview_rate"]}% reach interview ({best["interviews"]}/{best["count"]}). '
                f'This beats your {metrics["interview_rate"]}% overall — prioritise more of these.',
                f'Search & apply to more {best["segment"]} roles')
        worst = [s for s in fam if s["interview_rate"]==0]
        if worst:
            w = max(worst, key=lambda x:x["count"])
            add(3,"strategy",f'Re-think your {w["segment"]} approach',
                f'{w["count"]} {w["segment"]} applications, 0 interviews so far. Either tailor the resume/keywords for these or reallocate effort to higher-converting families.',
                "Review or pause this segment")

    # 4) Source insight (exclude low-value Weekday AI-screen mills)
    src = [s for s in metrics["by_source"] if s["count"]>=3 and "weekday" not in s["segment"].lower()]
    if src and src[0]["interview_rate"]>metrics["interview_rate"]:
        s=src[0]
        add(3,"insight",f'{s["segment"]} is your best channel',
            f'Roles via {s["segment"]} convert at {s["interview_rate"]}% vs {metrics["interview_rate"]}% overall. Favor postings on this channel.',
            None)

    # 5) Stale applications → re-engage or archive
    stale = [a for a in apps if a["stale"]]
    if stale:
        items=[{"label":f'{a["company"]} — {a["title"]}',"sub":f'{a["days_since"]}d, no reply'} for a in sorted(stale,key=lambda a:-(a["days_since"] or 0))[:10]]
        add(2,"cleanup",f'Re-engage or archive {len(stale)} stalled applications',
            "These have sat >21 days with no movement. Send the 're-engage' template to the recruiter, or mark them closed to keep your board honest.",
            "Use the Re-engage template", items)

    # 6) Networking queue — warm contacts not touched recently
    queue = networking_queue(store)
    if queue:
        items=[{"label":f'{c["name"]} — {c["company"]}',"sub":c["title"]} for c in queue[:8]]
        add(2,"network",f'Reach out to {len(queue)} warm contacts',
            "Relationships decay. These recruiters/managers already know you — a short check-in keeps you top of mind for current and future roles.",
            "Open Network tab → send a note", items)

    # 7) Momentum / cadence
    if metrics["applied_last_7"] < 5:
        add(3,"momentum","Keep your application momentum up",
            f'Only {metrics["applied_last_7"]} new applications in the last 7 days. Aim for a steady weekly cadence in your best-converting families to keep the pipeline full.',
            "Block time to apply this week")
    else:
        add(4,"momentum","Strong momentum 🔥",
            f'{metrics["applied_last_7"]} applications in the last 7 days. Nice pace — make sure follow-ups and prep keep up with volume.')

    # 7b) Job-alert leads → match to strong families
    leads = store.get("leads",[])
    fresh = [l for l in leads if (days_since(l.get("date","")) or 999) <= 10]
    if fresh:
        strong = {s["segment"] for s in metrics["by_family"] if s["count"]>=3 and s["interview_rate"]>=metrics["interview_rate"]}
        matched=[l for l in fresh if job_family(l.get("title",""))in strong]
        items=[{"label":f'{l.get("company","?")} — {l.get("title","")}',"sub":job_family(l.get("title",""))} for l in (matched or fresh)[:8]]
        add(3,"leads",f'{len(fresh)} new job leads from alerts'+(f' · {len(matched)} in your strong families' if matched else ''),
            "From your job-alert emails. Prioritise the ones in families where you already convert well — that's where your effort pays off.",
            "Apply to the best-fit leads", items)

    # 8) Interview prep / recent rounds
    ints = sorted(store["interviews"], key=lambda x:x.get("date",""), reverse=True)
    if ints:
        recent = ints[:3]
        items=[{"label":f'{i["company"]} — {i["round"]}',"sub":i["date"]} for i in recent]
        add(4,"prep","Recent interview activity",
            "Review notes from these rounds and confirm any outstanding next steps.",
            "Open Interviews tab", items)

    recs.sort(key=lambda r:r["priority"])
    return recs

def networking_queue(store):
    q=[]
    for c in store.get("contacts",[]):
        w = c.get("warmth","cold")
        lt = c.get("last_touch","")
        gap = days_since(lt) if lt else 999
        if w in ("hot","warm") and (gap is None or gap>=10):
            cc=dict(c); cc["days_since_touch"]=gap; q.append(cc)
    return sorted(q, key=lambda c:({"hot":0,"warm":1,"cold":2}[c.get("warmth","cold")], -(c.get("days_since_touch") or 0)))

# ────────────────────────────────────────────────────────── payload + crypto
def build_payload(store):
    apps = enrich(store)
    metrics = compute_metrics(apps)
    recs = recommendations(apps, store, metrics)
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "owner": store["meta"].get("owner_name","You"),
        "metrics": metrics,
        "recommendations": recs,
        "applications": sorted(apps, key=lambda a:a["id"]),
        "interviews": sorted(store["interviews"], key=lambda x:x.get("date",""), reverse=True),
        "contacts": store.get("contacts",[]),
        "networking_queue": networking_queue(store),
        "templates": store.get("templates",[]),
        "leads": sorted(store.get("leads",[]), key=lambda l:l.get("date",""), reverse=True)[:60],
    }

def get_passphrase():
    if os.path.exists(PASSF):
        return open(PASSF).read().strip()
    pw = "changeme-careerboard"          # default; user changes in Phase B
    open(PASSF,"w").write(pw); os.chmod(PASSF,0o600)
    return pw

def encrypt_payload(payload, passphrase):
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    salt = os.urandom(16); iv = os.urandom(12); iters = 200_000
    kdf = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=iters)
    key = kdf.derive(passphrase.encode())
    ct  = AESGCM(key).encrypt(iv, json.dumps(payload).encode(), None)
    b64 = lambda b: base64.b64encode(b).decode()
    return {"v":1,"kdf":"PBKDF2-SHA256","iter":iters,
            "salt":b64(salt),"iv":b64(iv),"ct":b64(ct)}

def cmd_build():
    store = json.load(open(STORE))
    os.makedirs(WEBDIR, exist_ok=True)
    payload = build_payload(store)
    json.dump(payload, open(PLAIN,"w"), indent=2)          # local plaintext (served on localhost)
    try:                                                    # encrypted file (for phone/Vercel)
        enc = encrypt_payload(payload, get_passphrase())
        json.dump(enc, open(ENC,"w"))
        enc_msg = f"   encrypted → {ENC}"
    except ImportError:
        enc_msg = "   (encryption skipped: 'cryptography' not installed — local link still works)"
    m = payload["metrics"]
    print(f"✅ build OK — {m['total']} apps, {m['interview_rate']}% interview rate, "
          f"{len(payload['recommendations'])} recommendations, "
          f"{len(payload['networking_queue'])} contacts queued.")
    print(f"   plaintext → {PLAIN}\n{enc_msg}")

def cmd_cursor():
    store = json.load(open(STORE))
    c = store.get("cursor",{})
    print(json.dumps({"last_scan_iso":c.get("last_scan_iso"),
                      "processed_count":len(c.get("processed_thread_ids",[]))}))

def cmd_ingest(path):
    store = json.load(open(STORE))
    new = json.load(open(path))
    seen_threads = set(store["cursor"].get("processed_thread_ids",[]))
    pairs = {_app_key(a["company"], a["title"], a.get("date_applied","")) for a in store["applications"]}
    next_id = max((a["id"] for a in store["applications"]), default=0)+1
    added_apps=added_int=updated=0
    for rec in new.get("applications",[]):
        tid=rec.get("thread_id")
        if tid and tid in seen_threads:
            continue
        key=_app_key(rec["company"], rec["title"], rec.get("date_applied",""))
        if key in pairs:
            # update stage if more advanced / rejection
            for a in store["applications"]:
                if _app_key(a["company"],a["title"],a.get("date_applied",""))==key:
                    if rec.get("stage") and rec["stage"]!=a["stage"]:
                        a["stage"]=rec["stage"]; updated+=1
                    if rec.get("interview"): a["interview"]=True
            if tid: seen_threads.add(tid)
            continue
        store["applications"].append({
            "id":next_id,"title":rec["title"],"company":rec["company"],"url":rec.get("url",""),
            "stage":rec.get("stage","Applied"),"date_applied":rec.get("date_applied","")[:10],
            "source":rec.get("source","Other"),"reply":rec.get("reply",True),
            "interview":rec.get("interview",False),"flag":rec.get("flag",False),
            "notes":rec.get("notes","Auto-added by scan")})
        pairs.add(key); next_id+=1; added_apps+=1
        if tid: seen_threads.add(tid)
    for rec in new.get("interviews",[]):
        tid=rec.get("thread_id")
        if tid and tid in seen_threads: continue
        nid=max((i["id"] for i in store["interviews"]),default=0)+1
        store["interviews"].append({"id":nid,"company":rec["company"],"title":rec.get("title",""),
            "date":rec.get("date","")[:10],"round":rec.get("round",""),
            "interviewer":rec.get("interviewer",""),"details":rec.get("details","")})
        added_int+=1
        if tid: seen_threads.add(tid)
    for c in new.get("contacts",[]):
        if not any(x.get("email")==c.get("email") and c.get("email") for x in store["contacts"]):
            store["contacts"].append(c)
    store.setdefault("leads",[])
    lead_pairs={(_norm(l.get("company","")),_norm(l.get("title",""))) for l in store["leads"]}
    added_leads=0
    for l in new.get("leads",[]):
        k=(_norm(l.get("company","")),_norm(l.get("title","")))
        if k in lead_pairs or not k[1]: continue
        store["leads"].append({"company":l.get("company",""),"title":l.get("title",""),
            "source":l.get("source","Job alert"),"date":l.get("date","")[:10],"url":l.get("url","")})
        lead_pairs.add(k); added_leads+=1
    if added_leads: print(f"   +{added_leads} job leads")
    store["cursor"]["processed_thread_ids"]=sorted(seen_threads)
    store["cursor"]["last_scan_iso"]=new.get("scanned_at", datetime.now(timezone.utc).isoformat())
    json.dump(store, open(STORE,"w"), indent=2)
    print(f"✅ ingest — +{added_apps} apps, +{added_int} interviews, {updated} stage updates")
    cmd_build()

def cmd_xlsx():
    """Render a clean Job Tracker.xlsx mirror from store.json (also repairs the
    corrupted sheet dimension by writing a fresh workbook)."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    store = json.load(open(STORE)); apps = enrich(store)
    metrics = compute_metrics(apps)
    # canonical xlsx lives in the (non-TCC) project dir; Desktop has a symlink to it
    XLSX = os.path.join(ROOT, "Job Tracker.xlsx")

    HFILL=PatternFill("solid",fgColor="1F4E79"); HF=Font(name="Arial",bold=True,color="FFFFFF",size=11)
    YEL=PatternFill("solid",fgColor="FFFF00"); DF=Font(name="Arial",size=10); BF=Font(name="Arial",size=10,bold=True)
    C=Alignment("center","center",wrap_text=True); L=Alignment("left","center",wrap_text=True)
    thin=Side(style="thin",color="CCCCCC"); BD=Border(thin,thin,thin,thin)
    STAGE_COLOR={"Rejected":"FF0000","Final Round":"2E7D32","Round 2":"2E7D32","Offer Pending":"2E7D32",
                 "Offer – Declining":"E65100","Interview":"1565C0","Phone Screen":"E65100"}

    wb=Workbook();
    def sheet(name, headers, widths):
        ws=wb.create_sheet(name)
        for i,h in enumerate(headers,1):
            c=ws.cell(1,i,h); c.fill=HFILL; c.font=HF; c.alignment=C; c.border=BD
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
        ws.row_dimensions[1].height=30; ws.freeze_panes="A2"; return ws

    wb.remove(wb.active)
    # Application Tracker
    a=sheet("Application Tracker",
        ["App #","Job Title","Company","Job Desc. (hyperlink)","Application  Stage","Date Applied",
         "Source","People to reach out to (Linkedin) ","Outreach Message","Reply Email Received","Interview Received?"],
        [6,40,24,20,18,13,16,26,26,18,16])
    for r,ap in enumerate(sorted(apps,key=lambda x:x["id"]),2):
        fill = YEL if ap.get("flag") else None
        vals=[ap["id"],ap["title"],ap["company"],ap.get("url",""),ap["stage"],ap["date_applied"],
              ap["source"],"","","Yes" if ap.get("reply") else "","Yes" if ap.get("interview") else "No"]
        for i,v in enumerate(vals,1):
            c=a.cell(r,i,v); c.font=BF if i==1 else DF; c.border=BD
            c.alignment=C if i in (1,5,6,7,10,11) else L
            if fill: c.fill=fill
            if i==5 and ap["stage"] in STAGE_COLOR:
                c.font=Font(name="Arial",size=10,bold=True,color=STAGE_COLOR[ap["stage"]])
            if i==11 and ap.get("interview"): c.font=Font(name="Arial",size=10,bold=True,color="2E7D32")
    # Interview Tracker
    it=sheet("Interview Tracker",
        ["Interview #","Company name","Job Title","Interview Date","Interview Round","Interviewer Details","Interview Details"],
        [11,24,38,14,30,52,58])
    for r,iv in enumerate(sorted(store["interviews"],key=lambda x:x.get("date","")),2):
        vals=[iv["id"],iv["company"],iv["title"],iv["date"],iv["round"],iv["interviewer"],iv["details"]]
        for i,v in enumerate(vals,1):
            c=it.cell(r,i,v); c.font=BF if i==1 else DF; c.alignment=C if i in (1,4) else L; c.border=BD
        it.row_dimensions[r].height=38
    # Contacts
    ct=sheet("Contacts",["Name","Title","Company","Email","Phone","Warmth","Last Touch","Context","LinkedIn"],
        [20,30,22,30,15,10,13,50,28])
    for r,c0 in enumerate(store.get("contacts",[]),2):
        vals=[c0.get("name"),c0.get("title"),c0.get("company"),c0.get("email"),c0.get("phone",""),
              c0.get("warmth"),c0.get("last_touch",""),c0.get("context",""),c0.get("linkedin","")]
        for i,v in enumerate(vals,1):
            c=ct.cell(r,i,v); c.font=DF; c.alignment=L; c.border=BD
    # Outreach Templates
    tp=sheet("Outreach Templates",["Template","Audience","When to use","Message"],[34,16,40,90])
    for r,t in enumerate(store.get("templates",[]),2):
        vals=[t["name"],t["audience"],t["when"],t["body"]]
        for i,v in enumerate(vals,1):
            c=tp.cell(r,i,v); c.font=DF; c.alignment=L; c.border=BD
        tp.row_dimensions[r].height=120
    # Leads (from job alerts)
    lsheet=sheet("Leads (from alerts)",["Company","Job Title","Likely Family","Source","Date"],[24,40,22,18,14])
    for r,l in enumerate(sorted(store.get("leads",[]),key=lambda x:x.get("date",""),reverse=True),2):
        vals=[l.get("company"),l.get("title"),job_family(l.get("title","")),l.get("source"),l.get("date")]
        for i,v in enumerate(vals,1):
            c=lsheet.cell(r,i,v); c.font=DF; c.alignment=L if i in (1,2,3) else C; c.border=BD
    # Insights
    ins=sheet("Insights",["Metric","Value"],[34,16])
    rows=[("Total applications",metrics["total"]),("Reached interview",metrics["interviews"]),
          ("Interview rate",f'{metrics["interview_rate"]}%'),("Rejected",metrics["rejected"]),
          ("Still active",metrics["active"]),("Applied last 7 days",metrics["applied_last_7"]),
          ("Applied last 30 days",metrics["applied_last_30"]),("",""),("Interview rate by family","")]
    for fam in metrics["by_family"]: rows.append((f'  {fam["segment"]}',f'{fam["interview_rate"]}% ({fam["interviews"]}/{fam["count"]})'))
    rows.append(("","")); rows.append(("Interview rate by source",""))
    for s in metrics["by_source"]:
        if s["count"]>=2: rows.append((f'  {s["segment"]}',f'{s["interview_rate"]}% ({s["interviews"]}/{s["count"]})'))
    for r,(k,v) in enumerate(rows,2):
        ck=ins.cell(r,1,k); cv=ins.cell(r,2,v)
        ck.font=BF if (k and not k.startswith("  ")) else DF; cv.font=DF
        ck.alignment=L; cv.alignment=L; ck.border=BD; cv.border=BD
    # Config
    cf=sheet("Config",["Key","Value"],[22,28])
    for r,(k,v) in enumerate([("last_scan_iso",store["cursor"].get("last_scan_iso")),
            ("processed_threads",len(store["cursor"].get("processed_thread_ids",[]))),
            ("total_applications",metrics["total"]),("total_interviews",len(store["interviews"])),
            ("dashboard","http://localhost:8765")],2):
        ck=cf.cell(r,1,k); cv=cf.cell(r,2,v); ck.font=DF; cv.font=DF; ck.border=BD; cv.border=BD
    wb.save(XLSX)
    print(f"✅ xlsx mirror written (repaired) → {XLSX}")

if __name__=="__main__":
    cmd = sys.argv[1] if len(sys.argv)>1 else "build"
    if cmd=="build": cmd_build()
    elif cmd=="cursor": cmd_cursor()
    elif cmd=="ingest": cmd_ingest(sys.argv[2])
    elif cmd=="xlsx": cmd_xlsx()
    else: print("usage: career_lib.py [build|cursor|ingest <file>|xlsx]"); sys.exit(1)
